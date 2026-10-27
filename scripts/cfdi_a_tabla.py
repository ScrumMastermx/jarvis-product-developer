#!/usr/bin/env python3
"""Convierte una carpeta de CFDI 4.0 (XML del SAT) en un Excel con hojas Facturas, Resumen, Resumen emisor y Fuente.

Uso: python scripts/cfdi_a_tabla.py <carpeta_con_xml> <salida.xlsx>

Los XML que no se puedan leer o no sean CFDI 4.0 se reportan en pantalla (stderr) y no detienen el proceso.
Cuenta: XML en la carpeta = filas en Facturas + XML inválidos. Los UUID repetidos SÍ son filas (columna "Duplicado")
pero no entran al Resumen. Un CFDI sin UUID se marca "SIN UUID" y nunca se trata como duplicado.
Resumen: Ingresos (I), Egresos (E, notas de crédito) y Neto (I - E); los tipos P (pago), T (traslado) y N (nómina)
solo se cuentan, nunca se suman a los importes.
"""
import sys
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

NS_CFDI = "http://www.sat.gob.mx/cfd/4"
NS_TFD = "http://www.sat.gob.mx/TimbreFiscalDigital"
CENT = Decimal("0.01")
CERO = Decimal(0)

COLUMNAS = [
    "UUID", "Fecha", "Tipo", "Serie", "Folio", "Emisor RFC", "Emisor nombre", "Receptor RFC",
    "Receptor nombre", "Uso CFDI", "Subtotal", "Descuento", "IVA trasladado", "OtrosTraslados", "Retenciones",
    "Total", "Moneda", "Forma de pago", "Método de pago", "Duplicado", "Cuadra", "Archivo",
]
DINERO = ("Subtotal", "Descuento", "IVA trasladado", "OtrosTraslados", "Retenciones", "Total")


def dec(valor) -> Decimal:
    """Decimal desde el texto del atributo (nunca desde float). Vacío o ausente = 0."""
    return Decimal(valor.strip()) if valor and valor.strip() else CERO


def leer_cfdi(path: Path) -> dict:
    root = ET.parse(path).getroot()
    if root.tag != f"{{{NS_CFDI}}}Comprobante" or root.get("Version") != "4.0":
        raise ValueError("no es un cfdi:Comprobante versión 4.0")
    emisor = root.find(f"{{{NS_CFDI}}}Emisor")
    receptor = root.find(f"{{{NS_CFDI}}}Receptor")
    timbre = root.find(f".//{{{NS_TFD}}}TimbreFiscalDigital")
    if emisor is None or receptor is None or timbre is None:
        raise ValueError("falta Emisor, Receptor o TimbreFiscalDigital")

    iva = otros = ret = CERO
    nodo = root.find(f"{{{NS_CFDI}}}Impuestos")
    if nodo is not None:
        todos = CERO
        for tras in nodo.iterfind(f"{{{NS_CFDI}}}Traslados/{{{NS_CFDI}}}Traslado"):
            imp = dec(tras.get("Importe"))
            todos += imp
            if tras.get("Impuesto") == "002":
                iva += imp
        total_tras = dec(nodo.get("TotalImpuestosTrasladados")) if nodo.get("TotalImpuestosTrasladados") else todos
        otros = total_tras - iva
        if nodo.get("TotalImpuestosRetenidos"):
            ret = dec(nodo.get("TotalImpuestosRetenidos"))
        else:
            ret = sum((dec(r.get("Importe")) for r in nodo.iterfind(f"{{{NS_CFDI}}}Retenciones/{{{NS_CFDI}}}Retencion")), CERO)

    subtotal, desc, total = dec(root.get("SubTotal")), dec(root.get("Descuento")), dec(root.get("Total"))
    cuadra = "" if abs(subtotal - desc + iva + otros - ret - total) <= CENT else "NO CUADRA"
    fecha = datetime.fromisoformat(root.get("Fecha"))
    return {
        "UUID": timbre.get("UUID", "").strip().upper() or "SIN UUID",
        "Fecha": fecha,
        "Tipo": root.get("TipoDeComprobante", ""),
        "Serie": root.get("Serie", ""),
        "Folio": root.get("Folio", ""),
        "Emisor RFC": emisor.get("Rfc", "").strip().upper(),
        "Emisor nombre": emisor.get("Nombre", ""),
        "Receptor RFC": receptor.get("Rfc", "").strip().upper(),
        "Receptor nombre": receptor.get("Nombre", ""),
        "Uso CFDI": receptor.get("UsoCFDI", ""),
        "Subtotal": subtotal,
        "Descuento": desc,
        "IVA trasladado": iva,
        "OtrosTraslados": otros,
        "Retenciones": ret,
        "Total": total,
        "Moneda": root.get("Moneda", ""),
        "Forma de pago": root.get("FormaPago", ""),
        "Método de pago": root.get("MetodoPago", ""),
        "Duplicado": "",
        "Cuadra": cuadra,
        "Archivo": path.name,
    }


def num(d: Decimal) -> float:
    return float(d.quantize(CENT))


def estilo_encabezado(ws, fila=1):
    for c in ws[fila]:
        c.font = Font(bold=True)


def escribir_tabla(ws, encabezados, filas):
    ws.append(encabezados)
    estilo_encabezado(ws)
    for fila in filas:
        ws.append(fila)
    ws.freeze_panes = "A2"
    for i, h in enumerate(encabezados, 1):
        ancho = max([len(str(h))] + [len(str(r[i - 1])) for r in filas if r[i - 1] is not None][:200])
        ws.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 10), 50)


def hoja_resumen(wb, titulo, claves, acumulado, etiquetas):
    """Una tabla: columnas de texto + cifras por tipo; fila TOTAL con =SUM; Neto = Ingresos - Egresos."""
    ws = wb.create_sheet(titulo)
    k = len(etiquetas)
    ws.append(list(etiquetas) + ["Facturas I", "Ingresos (I)", "Facturas E", "Egresos (E)", "Neto (I-E)",
                                 "IVA neto (I-E)", "Pagos P (cuenta)", "Traslados T (cuenta)", "Nómina N (cuenta)"])
    estilo_encabezado(ws)

    def col(i):  # i = 1..9, posición dentro de las cifras
        return get_column_letter(k + i)
    for clave in claves:
        a = acumulado[clave]
        r = ws.max_row + 1
        base = list(clave) if isinstance(clave, tuple) else [clave]
        ws.append(base + [a["I"][0], num(a["I"][1]), a["E"][0], num(a["E"][1]), f"={col(2)}{r}-{col(4)}{r}",
                          num(a["I"][2] - a["E"][2]), a["P"][0], a["T"][0], a["N"][0]])
    n = ws.max_row
    if n >= 2:
        ws.append(["TOTAL"] + [""] * (k - 1) + [f"=SUM({col(i)}2:{col(i)}{n})" for i in range(1, 10)])
        estilo_encabezado(ws, ws.max_row)
    for i in range(1, 10):
        letra = col(i)
        ws.column_dimensions[letra].width = 18
        for c in ws[letra][1:]:
            c.number_format = "0" if i in (1, 3, 7, 8, 9) else "#,##0.00"
    for i in range(1, k + 1):
        ws.column_dimensions[get_column_letter(i)].width = 40 if i == 2 else 18
    ws.freeze_panes = "A2"


def main(argv):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(errors="replace")
        except Exception:
            pass
    if len(argv) != 3:
        print(__doc__, file=sys.stderr)
        return 2
    carpeta, salida = Path(argv[1]), Path(argv[2])
    if not carpeta.is_dir():
        print(f"ERROR: {carpeta.name or carpeta} no es una carpeta", file=sys.stderr)
        return 2

    archivos = sorted(p for p in carpeta.rglob("*") if p.suffix.lower() == ".xml")
    facturas, invalidos = [], []
    for p in archivos:
        rel = p.relative_to(carpeta).as_posix()
        try:
            f = leer_cfdi(p)
            f["Archivo"] = rel
            facturas.append(f)
        except (ET.ParseError, ValueError, TypeError, OSError, InvalidOperation) as e:
            motivo = "no se pudo leer (es una carpeta o no hay permiso)" if isinstance(e, OSError) else str(e)
            invalidos.append((rel, motivo))
            print(f"XML inválido: {rel}: {motivo}", file=sys.stderr)

    facturas.sort(key=lambda f: (f["Fecha"], f["UUID"], f["Archivo"]))
    vistos, duplicados = set(), 0
    for f in facturas:
        if f["UUID"] == "SIN UUID":
            continue
        if f["UUID"] in vistos:
            f["Duplicado"] = "SÍ"
            duplicados += 1
        else:
            vistos.add(f["UUID"])
    sin_uuid = sum(1 for f in facturas if f["UUID"] == "SIN UUID")
    no_cuadran = sum(1 for f in facturas if f["Cuadra"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Facturas"
    filas = [[num(f[c]) if c in DINERO else f[c] for c in COLUMNAS] for f in facturas]
    escribir_tabla(ws, COLUMNAS, filas)
    for col in DINERO:
        for celda in ws[get_column_letter(COLUMNAS.index(col) + 1)][1:]:
            celda.number_format = "#,##0.00"
    for celda in ws[get_column_letter(COLUMNAS.index("Fecha") + 1)][1:]:
        celda.number_format = "yyyy-mm-dd hh:mm"

    # Resumen: solo facturas no duplicadas. Importes solo para I y E; P/T/N solo cuenta.
    def nuevo():
        return {t: [0, CERO, CERO] for t in "IEPTN"}
    por_mes, por_emisor = defaultdict(nuevo), defaultdict(nuevo)
    otros_tipos = 0
    for f in facturas:
        if f["Duplicado"]:
            continue
        t = f["Tipo"]
        if t not in ("I", "E", "P", "T", "N"):
            otros_tipos += 1
            continue
        for clave, bucket in ((f["Fecha"].strftime("%Y-%m"), por_mes), ((f["Emisor RFC"], f["Emisor nombre"]), por_emisor)):
            b = bucket[clave][t]
            b[0] += 1
            if t in ("I", "E"):
                b[1] += f["Total"]
                b[2] += f["IVA trasladado"]
    hoja_resumen(wb, "Resumen", sorted(por_mes), por_mes, ["Mes"])
    hoja_resumen(wb, "Resumen emisor",
                 sorted(por_emisor, key=lambda k: (-(por_emisor[k]["I"][1] - por_emisor[k]["E"][1]), k)),
                 por_emisor, ["Emisor RFC", "Emisor nombre"])

    monedas = sorted({f["Moneda"] for f in facturas if f["Tipo"] in ("I", "E") and not f["Duplicado"]})
    fu = wb.create_sheet("Fuente")
    fu.append(["Dato", "Valor"])
    estilo_encabezado(fu)
    fu.append(["Carpeta", carpeta.resolve().name])
    fu.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    fu.append(["XML en la carpeta", len(archivos)])
    fu.append(["Filas en Facturas", len(facturas)])
    fu.append(["XML inválidos", len(invalidos)])
    fu.append(["Comprobación: XML = filas + inválidos", "OK" if len(archivos) == len(facturas) + len(invalidos) else "NO CUADRA"])
    fu.append(["Duplicados (UUID repetido; son filas, fuera del Resumen)", duplicados])
    fu.append(["Sin UUID (fila incluida, marcada SIN UUID)", sin_uuid])
    fu.append(["Filas donde Subtotal - Descuento + IVA + OtrosTraslados - Retenciones no es Total", no_cuadran])
    fu.append(["Tipos fuera de I/E/P/T/N (no entran al Resumen)", otros_tipos])
    fu.append(["Monedas en ingresos/egresos (el Resumen no convierte)", ", ".join(m or "(vacía)" for m in monedas)])
    for nombre, motivo in invalidos:
        fu.append(["Inválido: " + nombre, motivo])
    fu.column_dimensions["A"].width = 70
    fu.column_dimensions["B"].width = 60

    try:
        salida.parent.mkdir(parents=True, exist_ok=True)
        wb.save(salida)
    except OSError as e:
        print(f"ERROR: no se pudo escribir {salida.name}: {e.strerror or e} (¿está abierto en Excel?)", file=sys.stderr)
        return 1
    print(f"{len(facturas)} facturas ({duplicados} duplicadas, {sin_uuid} sin UUID, {no_cuadran} no cuadran), "
          f"{len(invalidos)} XML inválidos -> {salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
