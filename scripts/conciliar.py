#!/usr/bin/env python3
"""Conciliación bancaria: estado de cuenta (banco) vs libro/auxiliar contable.

Uso: python scripts/conciliar.py <banco.csv> <libro.csv> <salida.xlsx> [--dias N] [--invertir-libro]

banco.csv: fecha,descripcion,monto     libro.csv: fecha,referencia,monto
Fechas: AAAA-MM-DD, AAAA-MM-DD HH:MM:SS o DD/MM/AAAA. Montos con signo: cargos negativos, abonos positivos
(también acepta (350.00) = -350.00, $1,234.50 y celdas vacías, que se omiten con aviso).
Empareja por monto exacto (2 decimales) y fecha dentro de ±N días (3 por defecto): primero las parejas con
la misma fecha, después las de menor diferencia de días. Cada movimiento se usa una vez.
--invertir-libro cambia el signo de todo el libro (para libros que registran cargos como positivos).
La codificación del CSV se detecta sola (UTF-8, Windows-1252 o Latin-1) y se anota en la hoja Fuente.
"""
import argparse
import csv
import io
import re
import sys
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

CODIFICACIONES = ("utf-8-sig", "cp1252", "latin-1")


class ErrorUsuario(Exception):
    """Error que se muestra en una sola línea en español."""


class Parser(argparse.ArgumentParser):
    def error(self, message):
        print(f"ERROR: argumentos incorrectos ({message}). Uso: python scripts/conciliar.py <banco.csv> <libro.csv> "
              f"<salida.xlsx> [--dias N] [--invertir-libro]", file=sys.stderr)
        sys.exit(1)


def sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def parse_monto(s: str) -> Decimal:
    """Devuelve el monto con 2 decimales. Lanza ValueError si no se entiende. Cadena vacía = ValueError('vacío')."""
    t = re.sub(r"[\s$ ]|MXN|USD", "", s.strip(), flags=re.I)
    if not t:
        raise ValueError("vacío")
    neg = False
    if t.startswith("(") and t.endswith(")"):
        neg, t = True, t[1:-1]
    if t.endswith("-"):
        neg, t = True, t[:-1]
    if t.startswith("-"):
        neg, t = True, t[1:]
    elif t.startswith("+"):
        t = t[1:]
    if "," in t and "." in t:
        # el último separador es el decimal
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(",", ".") if re.search(r",\d{1,2}$", t) and t.count(",") == 1 else t.replace(",", "")
    try:
        d = Decimal(t).quantize(Decimal("0.01"))
    except InvalidOperation:
        raise ValueError(f"monto no reconocido: {s.strip()!r}")
    return -d if neg else d


def parse_fecha(s: str, orden_dm: bool) -> datetime:
    s = s.strip()
    m = re.fullmatch(r"(\d{4})[-/](\d{1,2})[-/](\d{1,2})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?)?", s)
    if m:
        y, mo, d = (int(x) for x in m.groups())
    else:
        m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})(?:[ T]\d{1,2}:\d{2}(?::\d{2})?(?:\.\d+)?)?", s)
        if not m:
            raise ValueError(f"fecha no reconocida: {s!r} (usa AAAA-MM-DD o DD/MM/AAAA)")
        a, b, y = (int(x) for x in m.groups())
        d, mo = (a, b) if orden_dm else (b, a)
    try:
        return datetime(y, mo, d)
    except ValueError:
        raise ValueError(f"fecha inválida: {s!r}")


def decidir_orden(fechas_crudas):
    """Devuelve (orden_dm, nota). Con evidencia >12 decide; sin ella asume DD/MM y lo avisa."""
    pares = []
    for s in fechas_crudas:
        m = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-]\d{4}.*", s.strip())
        if m:
            pares.append((int(m.group(1)), int(m.group(2))))
    if not pares:
        return True, None
    primero_grande = any(a > 12 for a, _ in pares)
    segundo_grande = any(b > 12 for _, b in pares)
    if primero_grande and not segundo_grande:
        return True, None
    if segundo_grande and not primero_grande:
        return False, "las fechas con diagonal se leyeron como MM/DD/AAAA (hay días mayores a 12 en el segundo lugar)"
    if all(a <= 12 and b <= 12 for a, b in pares) and any(a != b for a, b in pares):
        return True, "las fechas con diagonal son ambiguas; se asumió DD/MM/AAAA (día/mes)"
    return True, None


def leer(path: Path, col_texto: str):
    """Devuelve (filas, omitidas, codificacion, notas)."""
    try:
        crudo = path.read_bytes()
    except OSError as e:
        raise ErrorUsuario(f"no se pudo abrir {path.name}: {e.strerror or e}")
    texto, usada = None, None
    for enc in CODIFICACIONES:
        try:
            texto, usada = crudo.decode(enc), enc
            break
        except UnicodeDecodeError:
            continue
    primera = texto.splitlines()[0] if texto and texto.splitlines() else ""
    delim = ";" if primera.count(";") > primera.count(",") else ","
    rd = csv.DictReader(io.StringIO(texto, newline=""), delimiter=delim)
    cab = {sin_acentos(h).strip().lower(): h for h in (rd.fieldnames or []) if h}
    faltan = {"fecha", col_texto, "monto"} - set(cab)
    if faltan:
        raise ErrorUsuario(f"a {path.name} le faltan columnas: {', '.join(sorted(faltan))} (encabezados leídos: {', '.join(cab) or 'ninguno'})")
    crudas = []
    for n, row in enumerate(rd, start=2):
        v = {k: (row.get(cab[k]) or "").strip() for k in ("fecha", col_texto, "monto")}
        if any(v.values()):
            crudas.append((n, v))
    orden_dm, nota = decidir_orden([v["fecha"] for _, v in crudas])
    notas = [f"{path.name}: {nota}"] if nota else []
    filas, omitidas = [], []
    for n, v in crudas:
        try:
            monto = parse_monto(v["monto"])
        except ValueError as e:
            if str(e) == "vacío":
                omitidas.append((n, "monto vacío"))
                continue
            raise ErrorUsuario(f"{path.name}, fila {n}: {e}")
        try:
            fecha = parse_fecha(v["fecha"], orden_dm)
        except ValueError as e:
            raise ErrorUsuario(f"{path.name}, fila {n}: {e}")
        filas.append({"fila": n, "fecha": fecha, "texto": v[col_texto], "monto": monto})
    return filas, omitidas, usada, notas


def emparejar(banco, libro, dias):
    """Devuelve lista de (banco, libro). Pares de misma fecha primero, luego menor |Δdías|; orden estable por fila."""
    por_monto = {}
    for j, l in enumerate(libro):
        por_monto.setdefault(l["monto"], []).append(j)
    cand = []
    for i, b in enumerate(banco):
        for j in por_monto.get(b["monto"], ()):
            delta = abs((libro[j]["fecha"] - b["fecha"]).days)
            if delta <= dias:
                cand.append((delta, i, j))
    cand.sort()
    usados_b, usados_l, pares = set(), set(), []
    for _, i, j in cand:
        if i in usados_b or j in usados_l:
            continue
        usados_b.add(i)
        usados_l.add(j)
        pares.append((i, j))
    pares.sort()
    return pares


def escribir(ws, encabezados, filas):
    ws.append(encabezados)
    for c in ws[1]:
        c.font = Font(bold=True)
    for f in filas:
        ws.append(f)
    ws.freeze_panes = "A2"
    for i, h in enumerate(encabezados, 1):
        letra = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letra].width = 40 if h.startswith(("Descripción", "Referencia")) else 14
        for c in ws[letra][1:]:
            if h.startswith("Fecha"):
                c.number_format = "yyyy-mm-dd"
            elif h == "Monto":
                c.number_format = "#,##0.00"
            elif h.startswith("Fila"):
                c.number_format = "0"


def construir(args, banco, libro, om_b, om_l, enc_b, enc_l, notas, invertir):
    pares = emparejar(banco, libro, args.dias)
    ub, ul = {i for i, _ in pares}, {j for _, j in pares}
    conciliadas = [[banco[i]["fecha"], banco[i]["texto"], float(banco[i]["monto"]), banco[i]["fila"],
                    libro[j]["fecha"], libro[j]["texto"], libro[j]["fila"]] for i, j in pares]
    solo_banco = [[b["fecha"], b["texto"], float(b["monto"]), b["fila"]] for i, b in enumerate(banco) if i not in ub]
    solo_libro = [[l["fecha"], l["texto"], float(l["monto"]), l["fila"]] for j, l in enumerate(libro) if j not in ul]

    wb = Workbook()
    ws = wb.active
    ws.title = "Conciliadas"
    escribir(ws, ["Fecha banco", "Descripción banco", "Monto", "Fila banco", "Fecha libro", "Referencia libro", "Fila libro"], conciliadas)
    escribir(wb.create_sheet("Solo banco"), ["Fecha", "Descripción", "Monto", "Fila"], solo_banco)
    escribir(wb.create_sheet("Solo libro"), ["Fecha", "Referencia", "Monto", "Fila"], solo_libro)

    def rango(hoja, n):
        return f"'{hoja}'!C2:C{max(n, 1) + 1}"
    rs = wb.create_sheet("Resumen")
    rs.append(["Concepto", "Movimientos", "Suma"])
    rs.append(["Conciliadas", f"=COUNT({rango('Conciliadas', len(conciliadas))})", f"=SUM({rango('Conciliadas', len(conciliadas))})"])
    rs.append(["Solo banco", f"=COUNT({rango('Solo banco', len(solo_banco))})", f"=SUM({rango('Solo banco', len(solo_banco))})"])
    rs.append(["Solo libro", f"=COUNT({rango('Solo libro', len(solo_libro))})", f"=SUM({rango('Solo libro', len(solo_libro))})"])
    rs.append(["Banco (total) = conciliadas + solo banco", "=B2+B3", "=C2+C3"])
    rs.append(["Libro (total) = conciliadas + solo libro", "=B2+B4", "=C2+C4"])
    rs.append(["Diferencia banco - libro", "", "=C5-C6"])
    rs.append(["Comprobación: solo banco - solo libro (debe ser igual a la diferencia)", "", "=C3-C4"])
    rs.append(["Control: banco leído del archivo", len(banco), float(sum(b["monto"] for b in banco))])
    rs.append(["Control: libro leído del archivo", len(libro), float(sum(l["monto"] for l in libro))])
    rs.append(["Control: diferencia contra las hojas (debe ser 0)", "=B9-B5+B10-B6", "=ROUND(C9-C5+C10-C6,2)"])
    for c in rs[1]:
        c.font = Font(bold=True)
    rs.column_dimensions["A"].width = 70
    rs.column_dimensions["B"].width = 14
    rs.column_dimensions["C"].width = 16
    for fila in rs.iter_rows(min_row=2):
        fila[1].number_format = "0"
        fila[2].number_format = "#,##0.00"

    fu = wb.create_sheet("Fuente")
    fu.append(["Dato", "Valor"])
    for c in fu[1]:
        c.font = Font(bold=True)
    fu.append(["Archivo banco", Path(args.banco).name])
    fu.append(["Archivo libro", Path(args.libro).name])
    fu.append(["Codificación banco", enc_b])
    fu.append(["Codificación libro", enc_l])
    fu.append(["Filas leídas banco", len(banco)])
    fu.append(["Filas leídas libro", len(libro)])
    fu.append(["Criterio", f"monto exacto y fecha ±{args.dias} días"])
    fu.append(["Ventana (días)", args.dias])
    fu.append(["Signo del libro invertido", "sí" if invertir else "no"])
    fu.append(["Generado", datetime.now().strftime("%Y-%m-%d %H:%M")])
    for n in notas:
        fu.append(["Nota", n])
    for nombre, om in (("banco", om_b), ("libro", om_l)):
        for fila, motivo in om:
            fu.append([f"Fila omitida ({nombre}, fila {fila})", motivo])
    fu.column_dimensions["A"].width = 40
    fu.column_dimensions["B"].width = 70
    return wb, conciliadas, solo_banco, solo_libro


def main(argv):
    for s in (sys.stdout, sys.stderr):
        try:
            s.reconfigure(errors="replace")
        except Exception:
            pass
    p = Parser(description="Conciliación bancaria banco vs libro", add_help=True)
    p.add_argument("banco")
    p.add_argument("libro")
    p.add_argument("salida")
    p.add_argument("--dias", type=int, default=3)
    p.add_argument("--invertir-libro", action="store_true")
    args = p.parse_args(argv[1:])
    if args.dias < 0:
        p.error("--dias debe ser 0 o más")
    try:
        banco, om_b, enc_b, nb = leer(Path(args.banco), "descripcion")
        libro, om_l, enc_l, nl = leer(Path(args.libro), "referencia")
        for nombre, om in ((Path(args.banco).name, om_b), (Path(args.libro).name, om_l)):
            for fila, motivo in om:
                print(f"AVISO: {nombre}, fila {fila}: {motivo}; se omite", file=sys.stderr)
        notas = nb + nl
        for n in notas:
            print(f"AVISO: {n}", file=sys.stderr)
        invertir = args.invertir_libro
        if invertir:
            for l in libro:
                l["monto"] = -l["monto"]
        elif banco and libro and not emparejar(banco, libro, args.dias):
            invertido = [dict(l, monto=-l["monto"]) for l in libro]
            if emparejar(banco, invertido, args.dias):
                print("AVISO: 0 conciliadas, pero con el signo del libro invertido sí hay coincidencias. "
                      "Si el libro registra cargos como positivos, vuelve a correr con --invertir-libro.", file=sys.stderr)
        print(f"Codificación: banco {enc_b}, libro {enc_l}")
        wb, conc, sb, sl = construir(args, banco, libro, om_b, om_l, enc_b, enc_l, notas, invertir)
        salida = Path(args.salida)
        salida.parent.mkdir(parents=True, exist_ok=True)
        wb.save(salida)
    except ErrorUsuario as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1
    except OSError as e:
        print(f"ERROR: no se pudo escribir {Path(args.salida).name}: {e.strerror or e} (¿está abierto en Excel?)", file=sys.stderr)
        return 1
    print(f"{len(conc)} conciliadas, {len(sb)} solo banco, {len(sl)} solo libro -> {args.salida}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
